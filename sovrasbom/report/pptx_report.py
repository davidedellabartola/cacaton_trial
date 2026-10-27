"""Generatore della presentazione PowerPoint (.pptx) con python-pptx.

Scelte di progetto:
- grafici NATIVI di PowerPoint (non immagini): restano modificabili e i dati sono ispezionabili
- una slide = un messaggio; ogni slide ha note per chi presenta
- le slide che non hanno dati (es. nessuna SBOM del fornitore) vengono saltate
"""

from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION, XL_LEGEND_POSITION
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Emu, Inches, Pt

from ..model import SEVERITY_IT, SEVERITY_ORDER
from . import style as S
from .markdown_report import online_note


# ============================================================ helper di base
def _font(run, size=S.BODY_SIZE, color=S.TEXT, bold=False, face=S.BODY_FONT, italic=False):
    run.font.size = size
    run.font.color.rgb = color
    run.font.bold = bold
    run.font.italic = italic
    run.font.name = face


def text_box(slide, x, y, w, h, text, size=S.BODY_SIZE, color=S.TEXT, bold=False,
             face=S.BODY_FONT, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, italic=False, name=None):
    """Casella di testo. `text` puo' essere una stringa o una lista di paragrafi."""
    tb = slide.shapes.add_textbox(x, y, w, h)
    if name:
        tb.name = name
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = 0
    tf.margin_top = tf.margin_bottom = 0
    paragraphs = text if isinstance(text, list) else [text]
    for i, para in enumerate(paragraphs):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        if i:
            p.space_before = Pt(6)
        r = p.add_run()
        r.text = str(para)
        _font(r, size, color, bold, face, italic)
    return tb


def rect(slide, x, y, w, h, fill, shape=MSO_SHAPE.ROUNDED_RECTANGLE, name=None):
    shp = slide.shapes.add_shape(shape, x, y, w, h)
    shp.fill.solid()
    shp.fill.fore_color.rgb = fill
    shp.line.fill.background()
    shp.shadow.inherit = False
    if shape == MSO_SHAPE.ROUNDED_RECTANGLE:
        shp.adjustments[0] = 0.08
    if name:
        shp.name = name
    return shp


def circle_label(slide, x, y, d, text, fill, size=Pt(14), color=S.WHITE):
    c = rect(slide, x, y, d, d, fill, shape=MSO_SHAPE.OVAL)
    tf = c.text_frame
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    r.text = str(text)
    _font(r, size, color, True)
    return c


def notes(slide, text):
    slide.notes_slide.notes_text_frame.text = text


def short(s, n):
    s = str(s or "")
    return s if len(s) <= n else s[: n - 1] + "…"


def pct(x):
    return f"{round(100 * x)}%"


# ===================================================== cornice delle slide
class Deck:
    def __init__(self, meta):
        self.prs = Presentation()
        self.prs.slide_width = S.SLIDE_W
        self.prs.slide_height = S.SLIDE_H
        self.blank = self.prs.slide_layouts[6]
        self.meta = meta
        self.page = 0

    def new_slide(self, title, kicker=""):
        """Slide di contenuto: titolo sempre nella stessa posizione, piede con numero pagina."""
        self.page += 1
        s = self.prs.slides.add_slide(self.blank)
        s.background.fill.solid()
        s.background.fill.fore_color.rgb = S.WHITE
        if kicker:
            text_box(s, Inches(0.6), Inches(0.38), Inches(12), Inches(0.3), kicker.upper(),
                     size=Pt(11), color=S.TEAL, bold=True, name="Kicker")
        text_box(s, Inches(0.6), Inches(0.68), Inches(12.1), Inches(0.75), title,
                 size=S.TITLE_SIZE, color=S.INK, bold=True, face=S.TITLE_FONT, name="Title")
        footer = f"SovraSBOM  ·  {self.meta.get('target', '')}  ·  {self.meta.get('date', '')}"
        text_box(s, Inches(0.6), Inches(7.02), Inches(10), Inches(0.25), footer,
                 size=S.CAPTION_SIZE, color=S.MUTED, name="Footer")
        text_box(s, Inches(11.73), Inches(7.02), Inches(1), Inches(0.25), str(self.page + 1),
                 size=S.CAPTION_SIZE, color=S.MUTED, align=PP_ALIGN.RIGHT, name="PageNumber")
        return s

    def dark_slide(self):
        s = self.prs.slides.add_slide(self.blank)
        s.background.fill.solid()
        s.background.fill.fore_color.rgb = S.INK
        return s


# ================================================================ grafici
def _style_chart(chart, font_size=Pt(11)):
    chart.font.name = S.BODY_FONT
    chart.font.size = font_size
    chart.font.color.rgb = S.MUTED


def bar_chart(slide, x, y, w, h, categories, values, colors, horizontal=False,
              number_format='0', max_value=None):
    """Grafico a barre nativo, un colore per barra, etichette con il valore."""
    if horizontal:   # nelle barre orizzontali PowerPoint disegna la prima categoria in basso
        categories, values, colors = categories[::-1], values[::-1], colors[::-1]
    data = CategoryChartData()
    data.categories = categories
    data.add_series("Valori", values)
    ctype = XL_CHART_TYPE.BAR_CLUSTERED if horizontal else XL_CHART_TYPE.COLUMN_CLUSTERED
    gframe = slide.shapes.add_chart(ctype, x, y, w, h, data)
    chart = gframe.chart
    _style_chart(chart)
    chart.has_legend = False
    chart.has_title = False
    plot = chart.plots[0]
    plot.gap_width = 55
    plot.vary_by_categories = False
    plot.has_data_labels = True
    dl = plot.data_labels
    dl.number_format = number_format
    dl.number_format_is_linked = False
    dl.position = XL_LABEL_POSITION.OUTSIDE_END
    dl.font.size = Pt(12)
    dl.font.bold = True
    dl.font.color.rgb = S.TEXT
    series = plot.series[0]
    for i, point in enumerate(series.points):
        point.format.fill.solid()
        point.format.fill.fore_color.rgb = colors[i % len(colors)]
    va = chart.value_axis
    va.has_major_gridlines = True
    va.major_gridlines.format.line.color.rgb = S.GRID
    va.format.line.fill.background()
    va.tick_labels.font.size = Pt(10)
    va.tick_labels.font.color.rgb = S.MUTED
    va.tick_labels.number_format = number_format
    va.tick_labels.number_format_is_linked = False
    va.minimum_scale = 0
    if max_value is not None:
        va.maximum_scale = max_value
    ca = chart.category_axis
    ca.format.line.color.rgb = S.LIGHT
    ca.tick_labels.font.size = Pt(11)
    ca.tick_labels.font.color.rgb = S.TEXT
    ca.has_major_gridlines = False
    return chart


def doughnut_chart(slide, x, y, w, h, categories, values, colors):
    data = CategoryChartData()
    data.categories = categories
    data.add_series("Componenti", values)
    chart = slide.shapes.add_chart(XL_CHART_TYPE.DOUGHNUT, x, y, w, h, data).chart
    _style_chart(chart)
    chart.has_title = False
    chart.has_legend = True
    chart.legend.position = XL_LEGEND_POSITION.RIGHT
    chart.legend.include_in_layout = False
    chart.legend.font.size = Pt(13)
    chart.legend.font.color.rgb = S.TEXT
    plot = chart.plots[0]
    plot.has_data_labels = False   # i valori sono nella legenda: leggibili su qualunque colore
    for i, point in enumerate(plot.series[0].points):
        point.format.fill.solid()
        point.format.fill.fore_color.rgb = colors[i]
        point.format.line.color.rgb = S.WHITE
    return chart


# ================================================================ tabelle
def table(slide, x, y, w, rows, col_widths, header, row_h=Inches(0.36), font=Pt(11),
          cell_colors=None):
    """Tabella con intestazione scura e righe alternate. cell_colors: {(r, c): RGBColor} per il testo."""
    n_rows = len(rows) + 1
    shape = slide.shapes.add_table(n_rows, len(header), x, y, w, row_h * n_rows)
    tbl = shape.table
    total = sum(col_widths)
    for i, cw in enumerate(col_widths):
        tbl.columns[i].width = Emu(int(w * cw / total))
    for r in range(n_rows):
        tbl.rows[r].height = row_h
        values = header if r == 0 else rows[r - 1]
        for c, val in enumerate(values):
            cell = tbl.cell(r, c)
            cell.margin_left = cell.margin_right = Inches(0.08)
            cell.margin_top = cell.margin_bottom = Inches(0.03)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            cell.fill.solid()
            if r == 0:
                cell.fill.fore_color.rgb = S.INK
            else:
                cell.fill.fore_color.rgb = S.TINT if r % 2 == 0 else S.WHITE
            tf = cell.text_frame
            tf.word_wrap = True
            p = tf.paragraphs[0]
            p.text = ""
            run = p.add_run()
            run.text = str(val)
            color = S.WHITE if r == 0 else (cell_colors or {}).get((r - 1, c), S.TEXT)
            _font(run, font, color, bold=(r == 0 or (cell_colors or {}).get((r - 1, c)) is not None))
    return tbl


def stat_card(slide, x, y, w, h, value, label, color=S.INK, sub=""):
    rect(slide, x, y, w, h, S.TINT)
    text_box(slide, x + Inches(0.25), y + Inches(0.18), w - Inches(0.5), Inches(0.75), str(value),
             size=Pt(40), color=color, bold=True, face=S.TITLE_FONT)
    text_box(slide, x + Inches(0.25), y + Inches(0.98), w - Inches(0.5), Inches(0.3), label,
             size=Pt(13), color=S.TEXT, bold=True)
    if sub:
        text_box(slide, x + Inches(0.25), y + Inches(1.28), w - Inches(0.5), Inches(0.3), sub,
                 size=S.SMALL_SIZE, color=S.MUTED)


def empty_state(slide, x, y, w, h, title, body, color=S.TEAL):
    rect(slide, x, y, w, h, S.TINT)
    circle_label(slide, x + Inches(0.35), y + Inches(0.35), Inches(0.55), "✓" if color == S.TEAL else "!", color,
                 size=Pt(18))
    text_box(slide, x + Inches(1.15), y + Inches(0.35), w - Inches(1.5), Inches(0.4), title,
             size=Pt(18), color=S.INK, bold=True)
    text_box(slide, x + Inches(1.15), y + Inches(0.85), w - Inches(1.5), h - Inches(1.1), body,
             size=S.BODY_SIZE, color=S.TEXT)


# ================================================================= slide
def _sov_color(sov):
    """Senza abbastanza dimensioni l'indice non e' un giudizio: niente rosso/verde."""
    return S.level_color(sov["index"]) if sov.get("classified", True) else S.MUTED


def _dims_text(sov):
    n = sov["measured"]
    return f"{n} {'dimensione' if n == 1 else 'dimensioni'} su {sov['possible']}"


def slide_title(deck, f, meta):
    s = deck.dark_slide()
    text_box(s, Inches(0.8), Inches(0.9), Inches(11), Inches(0.4),
             "ANALISI DELLA SUPPLY CHAIN SOFTWARE", size=Pt(14), color=S.LIGHT, bold=True)
    tlen = len(meta["target"])
    tsize = Pt(46) if tlen <= 32 else Pt(36) if tlen <= 55 else Pt(28)
    text_box(s, Inches(0.8), Inches(1.45), Inches(8.6), Inches(1.9), meta["target"],
             size=tsize, color=S.WHITE, bold=True, face=S.TITLE_FONT)
    text_box(s, Inches(0.8), Inches(3.5), Inches(9), Inches(0.9),
             "Provenienza, composizione e governance dei componenti: "
             "che cosa c'è davvero dentro il software",
             size=Pt(20), color=S.LIGHT, italic=True)
    git = meta.get("git") or {}
    origin = (f"Commit {git['commit'][:12]} del {git.get('commit_date', '')[:10]}" if git.get("commit")
              else short(meta.get("source", ""), 70))
    if meta.get("version"):
        origin += f"  ·  versione {meta['version']}"
    info = [x for x in [meta.get("organization"), f"Analisi del {meta['date']}", origin,
                        f"SBOM {meta['sbom_format']}  ·  {f['composition']['total']} componenti"] if x]
    text_box(s, Inches(0.8), Inches(4.85), Inches(8.4), Inches(1.6), info, size=Pt(14), color=S.LIGHT)
    sov = f["sovereignty"]
    col = _sov_color(sov)
    circle_label(s, Inches(10.15), Inches(1.6), Inches(2.3), sov["index"] if sov["measured"] else "–", col, size=Pt(60))
    text_box(s, Inches(9.65), Inches(4.05), Inches(3.3), Inches(0.35), "Indice di sovranità",
             size=Pt(14), color=S.WHITE, bold=True, align=PP_ALIGN.CENTER)
    sub = (f"NON CLASSIFICABILE · {_dims_text(sov)}" if not sov.get("classified", True)
           else f"PARZIALE · {_dims_text(sov)}" if sov.get("partial")
           else f"{sov['level']}  ·  su 100")
    text_box(s, Inches(9.35), Inches(4.42), Inches(3.9), Inches(0.35), sub,
             size=Pt(13), color=S.AMBER if sov.get("partial") else S.LIGHT, bold=bool(sov.get("partial")),
             align=PP_ALIGN.CENTER)
    text_box(s, Inches(0.8), Inches(6.85), Inches(11), Inches(0.3),
             f"Generato automaticamente da {meta['tool']} — evidenze e impronte SHA-256 allegate",
             size=S.CAPTION_SIZE, color=S.LIGHT)
    n_dims = sum(1 for v in sov["dimensions"].values() if v >= 0)
    notes(s, "Apertura: presentare il software analizzato e l'indice complessivo. "
             f"L'indice è la media di {n_dims} dimensioni spiegate nelle slide successive.")


def _key_messages(f):
    c, v, d, lic = f["composition"], f["vulns"], f["duplicates"], f["licenses"]
    pl, gv = f.get("pipeline") or {}, f.get("governance") or {}
    msgs = []
    if c["total"]:
        eco = len(c["by_ecosystem"])
        eco_txt = "1 ecosistema" if eco == 1 else f"{eco} ecosistemi diversi"
        msgs.append(f"Il software contiene {c['total']} componenti di {eco_txt}; "
                    f"{c['opaque']} sono binari o di tipo non identificato, i più difficili da aggiornare.")
    else:
        msgs.append("Nessun componente rilevato: il repository non dichiara dipendenze in formati leggibili "
                    "(o contiene solo la ricetta di build).")
    if v is not None and c["total"]:
        ch = v["by_severity"]["Critical"] + v["by_severity"]["High"]
        msgs.append(f"{v['total']} vulnerabilità note, di cui {ch} critiche o alte: "
                    f"{v['fixable']} hanno già una correzione disponibile"
                    + (f" e {len(v['kev'])} risultano sfruttate in attacchi reali." if v["kev"] else "."))
    if pl.get("applicable") and (pl["pipe_to_shell"] or pl["actions_third_party_unpinned"]):
        msgs.append(f"La build esegue {pl['pipe_to_shell']} volte codice scaricato senza verifica e usa "
                    f"{pl['actions_third_party_unpinned']} azioni di autori esterni non bloccate per impronta.")
    hist = gv.get("history") if gv.get("applicable") else None
    if hist and hist["top_author_share"] > 0.6:
        msgs.append(f"Un solo autore firma il {round(100 * hist['top_author_share'])}% degli ultimi "
                    f"{hist['commits']} commit: la continuità dipende da una persona.")
    if d["count"]:
        msgs.append(f"{d['count']} componenti compaiono in più versioni: copie incorporate che "
                    "l'aggiornamento del sistema operativo non corregge.")
    if c["total"]:
        msgs.append(f"Il {pct(lic['open_ratio'])} dei componenti ha una licenza open source riconosciuta; "
                    f"il {pct(lic['unknown_ratio'])} è senza licenza chiara.")
    return msgs[:3]


def slide_summary(deck, f, meta):
    s = deck.new_slide("Il quadro in sintesi", "Sintesi")
    v, sov = f["vulns"], f["sovereignty"]
    cards = [(f["composition"]["total"], "Componenti", S.INK, "rilevati nella SBOM")]
    if v is not None:
        ch = v["by_severity"]["Critical"] + v["by_severity"]["High"]
        cards.append((v["total"], "Vulnerabilità", S.INK, f"{ch} critiche o alte"))
        cards.append((len(v["kev"]), "Già sfruttate", S.RED if v["kev"] else S.TEAL, "catalogo CISA KEV"))
    else:
        cards.append(("—", "Vulnerabilità", S.MUTED, "scansione non eseguita"))
        pl = f.get("pipeline") or {}
        if pl.get("applicable"):
            high = sum(1 for x in pl["findings"] if x["severity"] == "Alta")
            cards.append((high, "Problemi gravi di build", S.RED if high else S.TEAL, "pipeline e Dockerfile"))
        else:
            cards.append((f["duplicates"]["count"], "Multi-versione", S.AMBER, "componenti duplicati"))
    cards.append((sov["index"], "Indice di sovranità", _sov_color(sov),
                  f"non classificabile: {_dims_text(sov)}" if not sov.get("classified", True)
                  else f"parziale: {sov['measured']} su {sov['possible']}" if sov.get("partial") else sov["level"]))
    w, gap = Inches(2.85), Inches(0.23)
    for i, (val, label, col, sub) in enumerate(cards):
        stat_card(s, Inches(0.6) + i * (w + gap), Inches(1.75), w, Inches(1.7), val, label, col, sub)
    text_box(s, Inches(0.6), Inches(3.85), Inches(6), Inches(0.4), "Messaggi chiave",
             size=Pt(20), color=S.INK, bold=True, face=S.TITLE_FONT)
    for i, msg in enumerate(_key_messages(f)):
        y = Inches(4.42) + i * Inches(0.82)
        circle_label(s, Inches(0.6), y, Inches(0.5), i + 1, S.TEAL)
        text_box(s, Inches(1.35), y + Inches(0.02), Inches(11.3), Inches(0.75), msg, size=S.BODY_SIZE)
    notes(s, "Questa slide da sola risponde alla domanda della sfida: quanto la PA conosce davvero "
             "ciò che c'è dentro il software. Le slide successive approfondiscono ogni numero.")


def slide_sovereignty(deck, f, meta):
    sov = f["sovereignty"]
    dims = [(k, v) for k, v in sov["dimensions"].items() if v >= 0]
    words = {1: "una dimensione misurabile", 2: "due dimensioni misurabili", 3: "tre dimensioni misurabili",
             4: "quattro dimensioni misurabili", 5: "cinque dimensioni misurabili", 6: "sei dimensioni misurabili",
             7: "sette dimensioni misurabili"}
    s = deck.new_slide(f"Indice di sovranità: {words.get(len(dims), 'dimensioni misurabili')}", "Sovranità")
    if not dims:
        empty_state(s, Inches(0.6), Inches(1.9), Inches(12.15), Inches(2.1), "Nessuna dimensione misurabile",
                    "La SBOM è vuota e l'input non è un repository di codice: non ci sono dati su cui calcolare "
                    "l'indice. Ripetere l'analisi con Syft o Trivy installati.", color=S.AMBER)
        return
    bar_chart(s, Inches(0.5), Inches(1.65), Inches(6.9), Inches(4.0),
              [k for k, _ in dims], [v for _, v in dims], [S.level_color(v) for _, v in dims],
              horizontal=True, max_value=100)
    expl = {
        "Trasparenza": "Quanto è completa la SBOM rispetto agli elementi minimi richiesti.",
        "Apertura": "Quota di componenti con licenza open source riconosciuta.",
        "Tracciabilità": "Componenti con identificativo univoco e origine nota.",
        "Sicurezza": "Penalità per vulnerabilità gravi, già sfruttate e diffuse.",
        "Manutenibilità": "Dipendenze ancora mantenute, non deprecate e aggiornate.",
        "Governance": "Regole, responsabili, continuità e strategia di uscita.",
        "Filiera di build": "Verificabilità del processo che produce il software.",
    }
    y = Inches(1.7)
    step = min(Inches(0.98), Emu(int(Inches(4.15) / max(1, len(dims)))))
    dense = len(dims) > 5
    for k, v in dims:
        size = Pt(14) if dense else Pt(16)
        text_box(s, Inches(7.85), y, Inches(3.9), Inches(0.3), k, size=size, color=S.INK, bold=True)
        text_box(s, Inches(11.75), y, Inches(1), Inches(0.3), f"{v}/100", size=size,
                 color=S.level_color(v), bold=True, align=PP_ALIGN.RIGHT)
        text_box(s, Inches(7.85), y + (Inches(0.27) if dense else Inches(0.34)), Inches(4.9), Inches(0.28),
                 expl.get(k, ""), size=Pt(11) if dense else S.SMALL_SIZE, color=S.MUTED)
        y += step
    rect(s, Inches(0.6), Inches(5.95), Inches(12.15), Inches(0.85), S.TINT)
    text_box(s, Inches(0.9), Inches(6.08), Inches(11.6), Inches(0.6),
             (f"Indice complessivo: {sov['index']}/100 ({sov['level']}). Media semplice delle dimensioni: "
              "la formula è pubblica e ricalcolabile dalle evidenze allegate." if sov.get("classified", True) else
              f"Media di {_dims_text(sov)}: {sov['index']}/100. Troppo poche per un giudizio complessivo sul "
              "software: il livello (buono, da migliorare, critico) richiede almeno tre dimensioni misurate."),
             size=S.BODY_SIZE, color=S.TEXT, anchor=MSO_ANCHOR.MIDDLE)
    notes(s, "Soglie: 75 o più = buono, 50-74 = da migliorare, sotto 50 = critico. "
             "Le formule sono descritte nella slide Metodologia.")


def slide_composition(deck, f, meta):
    c = f["composition"]
    s = deck.new_slide("Di cosa è fatto il software", "Composizione e provenienza")
    eco = c["by_ecosystem"][:8]
    colors = [S.INK if e not in ("binario", "sconosciuto") else S.AMBER for e, _ in eco]
    text_box(s, Inches(0.6), Inches(1.6), Inches(6.5), Inches(0.35), "Componenti per ecosistema",
             size=Pt(16), color=S.INK, bold=True)
    bar_chart(s, Inches(0.5), Inches(2.0), Inches(6.7), Inches(4.0),
              [e for e, _ in eco], [n for _, n in eco], colors, horizontal=True)
    text_box(s, Inches(0.6), Inches(6.15), Inches(6.5), Inches(0.6),
             "In arancione i componenti binari o non identificati: nessun gestore di pacchetti li aggiorna.",
             size=S.SMALL_SIZE, color=S.MUTED)
    text_box(s, Inches(7.6), Inches(1.6), Inches(5), Inches(0.35), "Principali origini dichiarate",
             size=Pt(16), color=S.INK, bold=True)
    rows = [(short(o, 34), n) for o, n in c["top_origins"][:7]]
    if rows:
        table(s, Inches(7.6), Inches(2.05), Inches(5.15), rows, [4, 1.2], ["Origine / fornitore", "Componenti"])
    y = Inches(2.05) + Inches(0.36) * (len(rows) + 1) + Inches(0.3)
    stat_card(s, Inches(7.6), y, Inches(2.45), Inches(1.6), c["opaque"], "Opachi",
              S.AMBER if c["opaque"] else S.TEAL, "binari / ignoti")
    stat_card(s, Inches(10.3), y, Inches(2.45), Inches(1.6), c["no_version"], "Senza versione",
              S.RED if c["no_version"] else S.TEAL, "non verificabili")
    notes(s, "L'origine è il namespace del purl (es. ubuntu) o il fornitore dichiarato. "
             "Un componente senza versione non può essere confrontato con i database di vulnerabilità.")


def slide_quality(deck, f, meta):
    q = f["quality"]
    s = deck.new_slide("Quanto ci si può fidare della SBOM", "Qualità dell'inventario")
    labels = list(q["coverage"].keys())
    vals = [round(100 * q["coverage"][k]) for k in labels]
    bar_chart(s, Inches(0.5), Inches(1.65), Inches(7.6), Inches(4.4), labels, vals,
              [S.level_color(v) for v in vals], horizontal=True, number_format='0"%"', max_value=100)
    text_box(s, Inches(0.6), Inches(6.2), Inches(7.4), Inches(0.6),
             "Percentuale di componenti che riportano ciascun elemento.",
             size=S.SMALL_SIZE, color=S.MUTED)
    rect(s, Inches(8.6), Inches(1.75), Inches(4.15), Inches(1.75), S.TINT)
    text_box(s, Inches(8.9), Inches(1.9), Inches(3.6), Inches(0.9), f"{q['score']}/100",
             size=Pt(44), color=S.level_color(q["score"]), bold=True, face=S.TITLE_FONT)
    text_box(s, Inches(8.9), Inches(2.85), Inches(3.6), Inches(0.5), "Punteggio elementi minimi NTIA",
             size=Pt(13), color=S.TEXT, bold=True)
    y = Inches(3.85)
    for label, ok in q["document"].items():
        circle_label(s, Inches(8.6), y, Inches(0.38), "✓" if ok else "✗", S.TEAL if ok else S.RED, size=Pt(12))
        text_box(s, Inches(9.15), y + Inches(0.03), Inches(3.6), Inches(0.38), label, size=Pt(14))
        y += Inches(0.6)
    text_box(s, Inches(8.6), y + Inches(0.1), Inches(4.15), Inches(0.9),
             "Elementi del documento: chi ha generato la SBOM, quando, e se descrive le dipendenze.",
             size=S.SMALL_SIZE, color=S.MUTED)
    notes(s, "Una SBOM incompleta dà falsa sicurezza: un componente senza versione o identificativo "
             "non viene trovato dai database di vulnerabilità. Licenza e hash non entrano nel punteggio NTIA "
             "ma sono mostrati perché servono alla sovranità e all'integrità.")


def slide_vulns(deck, f, meta):
    v = f["vulns"]
    s = deck.new_slide("Vulnerabilità note nei componenti", "Sicurezza")
    sev = [x for x in SEVERITY_ORDER if v["by_severity"][x] or x in ("Critical", "High", "Medium", "Low")]
    bar_chart(s, Inches(0.5), Inches(1.65), Inches(7.4), Inches(4.6),
              [SEVERITY_IT[x] for x in sev], [v["by_severity"][x] for x in sev],
              [S.SEVERITY_COLORS[x] for x in sev])
    text_box(s, Inches(0.6), Inches(6.3), Inches(7.2), Inches(0.5),
             f"Totale: {v['total']} vulnerabilità su {v['vulnerable_components']} componenti "
             f"({pct(v['component_ratio'])} dei componenti)."
             + (f" OSV.dev ha fornito i dettagli solo in parte: {v['details_missing']} advisory senza gravità."
                if v.get("details_missing") else ""), size=S.SMALL_SIZE,
             color=S.AMBER if v.get("details_missing") else S.MUTED)
    cards = [
        (v["fixable"], "Correggibili", S.AMBER if v["fixable"] else S.TEAL, "esiste già una versione corretta"),
        (len(v["kev"]), "Sfruttate (KEV)", S.RED if v["kev"] else S.TEAL, "usate in attacchi reali"),
        (v["epss_high"], "EPSS ≥ 10%", S.AMBER if v["epss_high"] else S.TEAL, "alta probabilità di attacco"),
    ]
    for i, (val, label, col, sub) in enumerate(cards):
        stat_card(s, Inches(8.4), Inches(1.75) + i * Inches(1.75), Inches(4.35), Inches(1.6), val, label, col, sub)
    notes(s, "Correggibili significa che il problema è già stato risolto a monte: il ritardo è del "
             "fornitore o della PA. KEV ed EPSS servono a decidere l'ordine degli interventi.")


def slide_priorities(deck, f, meta):
    v = f["vulns"]
    s = deck.new_slide("Cosa correggere per primo", "Priorità di intervento")
    rows, colors = [], {}
    for i, x in enumerate(v["prioritized"][:10]):
        fix = short(", ".join(x.fixed_in), 16) if x.fix_state == "fixed" and x.fixed_in else "nessuna"
        epss = f"{round(100 * x.epss, 1)}%" if x.epss >= 0 else "n/d"
        rows.append([x.id + (" *" if x.reach else ""), short(x.component, 22), short(x.version, 16),
                     SEVERITY_IT[x.severity], fix, "Sì" if x.kev else "No", epss])
        colors[(i, 3)] = S.MUTED if x.reach else S.SEVERITY_COLORS[x.severity]
        if x.kev:
            colors[(i, 5)] = S.RED
    if not rows:
        empty_state(s, Inches(0.6), Inches(1.9), Inches(12.15), Inches(2), "Nessuna vulnerabilità nota",
                    "La scansione non ha trovato vulnerabilità. Il risultato vale solo se il database "
                    "di Grype era aggiornato: verificarne la data nella slide Metodologia.")
        return
    table(s, Inches(0.6), Inches(1.7), Inches(12.15), rows, [2.2, 2.4, 1.7, 1.1, 1.7, 0.7, 0.9],
          ["Vulnerabilità", "Componente", "Versione", "Gravità", "Corretta in", "KEV", "EPSS"],
          row_h=Inches(0.4), font=Pt(12), cell_colors=colors)
    text_box(s, Inches(0.6), Inches(6.45), Inches(12), Inches(0.45),
             "Ordine: prima le vulnerabilità già sfruttate (KEV), poi impatto (CVSS) × probabilità (EPSS); "
             "a parità, sale chi ha già una correzione."
             + (" * Funzione vulnerabile non caricata o non chiamata dal progetto (libreria copiata): "
                "scende nell'ordine, ma va verificata." if any(x.reach for x in v["prioritized"]) else ""),
             size=S.SMALL_SIZE, color=S.MUTED)
    notes(s, "Stesso principio del rischio = probabilità × impatto. Mostrare che il tool non si limita a "
             "elencare CVE ma dice da dove cominciare.")


def slide_exposed(deck, f, meta):
    v = f["vulns"]
    if not v["top_components"]:
        return
    s = deck.new_slide("I componenti più esposti", "Sicurezza per componente")
    rows, colors = [], {}
    for i, c in enumerate(v["top_components"][:9]):
        rows.append([short(c["name"], 30), short(c["version"], 22), c["total"], c["crit_high"],
                     c["fixable"], "Sì" if c["kev"] else "No"])
        if c["crit_high"]:
            colors[(i, 3)] = S.RED
        if c["kev"]:
            colors[(i, 5)] = S.RED
    table(s, Inches(0.6), Inches(1.7), Inches(12.15), rows, [3.2, 2.2, 1.2, 1.4, 1.3, 0.8],
          ["Componente", "Versione", "Totale", "Critiche / alte", "Correggibili", "KEV"],
          row_h=Inches(0.4), font=Pt(12), cell_colors=colors)
    text_box(s, Inches(0.6), Inches(6.45), Inches(12), Inches(0.45),
             "Aggiornare pochi componenti spesso elimina la maggior parte del rischio.",
             size=S.SMALL_SIZE, color=S.MUTED)
    notes(s, "Utile per il dialogo con il fornitore: chiedere quando verranno aggiornati proprio questi componenti.")


def slide_licenses(deck, f, meta):
    lic = f["licenses"]
    s = deck.new_slide("Licenze: chi può modificare e mantenere il software", "Governance")
    if lic["unknown_ratio"] > 0.8 or lic.get("not_searched"):
        if lic.get("not_searched"):
            why = ("I file di dipendenze non riportano le licenze e l'analisi era offline, quindi il registro "
                   "(deps.dev) non è stato consultato: è un limite di questa analisi, non del software. "
                   "Ripeterla in rete o con Syft/Trivy.")
        else:
            why = ("Non è possibile dire se la PA può modificarli, riusarli o farli mantenere da altri. "
                   "Il dato va chiesto al fornitore o ricavato da un altro strumento.")
        empty_state(s, Inches(0.6), Inches(1.9), Inches(12.15), Inches(2.1), "Licenze non disponibili",
                    f"Il {pct(lic['unknown_ratio'])} dei componenti non ha una licenza nota. " + why +
                    " Per questo la dimensione Apertura dell'indice risulta non misurabile invece che pari a zero.",
                    color=S.AMBER)
        notes(s, "Un dato mancante non è un dato negativo: lo segnaliamo come tale invece di contarlo come zero.")
        return
    cats = [(k, n) for k, n in lic["by_category"].items() if n]
    if cats:
        doughnut_chart(s, Inches(0.4), Inches(1.6), Inches(7.0), Inches(4.7),
                       [f"{k}  ({n})" for k, n in cats], [n for _, n in cats], [S.LICENSE_COLORS[k] for k, _ in cats])
    text_box(s, Inches(7.85), Inches(1.6), Inches(4.9), Inches(0.35), "Licenze più frequenti",
             size=Pt(16), color=S.INK, bold=True)
    rows = [(short(l, 34), n) for l, n in lic["top_licenses"][:6]]
    if rows:
        table(s, Inches(7.85), Inches(2.05), Inches(4.9), rows, [4, 1.2], ["Licenza", "Componenti"])
    y = Inches(2.05) + Inches(0.36) * (len(rows) + 1) + Inches(0.25)
    compat = lic.get("compatibility") or []
    high = sum(len(x["components"]) for x in compat if x["severity"] == "Alta")
    other = sum(len(x["components"]) for x in compat if x["severity"] != "Alta")
    lines = [f"Open source riconosciuta: {pct(lic['open_ratio'])}",
             f"Senza licenza chiara: {pct(lic['unknown_ratio'])}",
             f"Proprietarie / source-available: {len(lic['proprietary'])}"]
    if lic.get("project_license"):
        lines.append(f"Compatibilità con {short(lic['project_license'], 14)}: "
                     + ("nessun problema" if not compat else f"{high} gravi, {other} da verificare"))
    rect(s, Inches(7.85), y, Inches(4.9), Inches(0.38 * len(lines) + 0.3), S.TINT)
    text_box(s, Inches(8.1), y + Inches(0.15), Inches(4.4), Inches(0.38 * len(lines)), lines,
             size=Pt(13), color=S.TEXT)
    if high:
        lines_y = y + Inches(0.38 * len(lines) + 0.4)
        text_box(s, Inches(7.85), lines_y, Inches(4.9), Inches(0.5),
                 short(next(x for x in compat if x["severity"] == "Alta")["message"], 110),
                 size=Pt(11), color=S.RED, bold=True)
    src = f" {lic['from_registry']} licenze sono state ricavate dal registro (deps.dev)." if lic.get("from_registry") else ""
    notes(s, "Una licenza aperta permette alla PA di far mantenere il software anche da altri fornitori "
             "(CAD artt. 68-69). 'Non riconosciuta' non significa proprietaria: va letta a mano." + src
             + "".join(f"\n[{x['severity']}] {x['message']}: {', '.join(x['components'][:8])}" for x in compat)
             + ("\nLa tabella di compatibilità è semplificata: non è consulenza legale." if compat else ""))


def slide_duplicates(deck, f, meta):
    d = f["duplicates"]
    s = deck.new_slide("Lo stesso componente, più versioni", "Librerie incorporate")
    if not d["count"]:
        empty_state(s, Inches(0.6), Inches(1.9), Inches(12.15), Inches(2), "Nessun componente duplicato",
                    "Ogni componente compare in una sola versione: gli aggiornamenti seguono un unico canale.")
        return
    rows = []
    for m in d["multi_version"][:8]:
        rows.append([short(m["name"], 24), short(", ".join(m["versions"]), 34), m["copies"],
                     short(", ".join(m["ecosystems"]), 18),
                     short(m["locations"][0] if m["locations"] else "—", 40)])
    table(s, Inches(0.6), Inches(1.7), Inches(12.15), rows, [2.2, 3.0, 0.9, 1.6, 3.6],
          ["Componente", "Versioni trovate", "Copie", "Ecosistemi", "Esempio di percorso"],
          row_h=Inches(0.42), font=Pt(12))
    text_box(s, Inches(0.6), Inches(6.35), Inches(12), Inches(0.55),
             "Le copie incorporate dal fornitore non vengono aggiornate dal sistema operativo: "
             "aggiornare il server non basta a correggerle.", size=S.SMALL_SIZE, color=S.MUTED)
    notes(s, "Nel caso Zimbra: OpenSSL e Java sotto /opt/zimbra/common sono copie separate da quelle di Ubuntu. "
             f"Totale componenti in più versioni: {d['count']}.")


def slide_vendor(deck, f, meta):
    vd = f["vendor_diff"]
    s = deck.new_slide("SBOM dichiarata dal fornitore e SBOM reale", "Verifica del fornitore")
    cards = [
        (pct(vd["coverage"]), "Copertura", S.level_color(round(100 * vd["coverage"])), "componenti reali dichiarati"),
        (len(vd["undeclared"]), "Non dichiarati", S.RED if vd["undeclared"] else S.TEAL, "presenti ma assenti nella SBOM"),
        (len(vd["mismatched"]), "Versioni diverse", S.AMBER if vd["mismatched"] else S.TEAL, "dichiarata ≠ installata"),
        (len(vd["not_found"]), "Non trovati", S.AMBER if vd["not_found"] else S.TEAL, "dichiarati ma assenti"),
    ]
    w, gap = Inches(2.85), Inches(0.23)
    for i, (val, label, col, sub) in enumerate(cards):
        stat_card(s, Inches(0.6) + i * (w + gap), Inches(1.7), w, Inches(1.65), val, label, col, sub)
    rows = [[short(n, 30), short(d, 22), short(o, 22)] for n, d, o in vd["mismatched"][:5]]
    if rows:
        text_box(s, Inches(0.6), Inches(3.65), Inches(6), Inches(0.35), "Versioni non corrispondenti",
                 size=Pt(16), color=S.INK, bold=True)
        table(s, Inches(0.6), Inches(4.08), Inches(7.2), rows, [3, 2.2, 2.2],
              ["Componente", "Dichiarata", "Rilevata"], font=Pt(12))
    und = [[short(n, 34)] for n in vd["undeclared"][:5]]
    if und:
        x = Inches(8.2) if rows else Inches(0.6)
        text_box(s, x, Inches(3.65), Inches(4.5), Inches(0.35), "Esempi di non dichiarati",
                 size=Pt(16), color=S.INK, bold=True)
        table(s, x, Inches(4.08), Inches(4.55), und, [1], ["Componente"], font=Pt(12))
    notes(s, f"SBOM del fornitore in formato {vd['vendor_fmt']} con {vd['declared']} componenti; "
             f"rilevati {vd['observed']} componenti sul sistema. Ogni differenza è una domanda precisa da porre al fornitore.")


def slide_provenance(deck, f, meta):
    pv = f["provenance"]
    s = deck.new_slide("Da dove arrivano davvero i componenti", "Provenienza opaca")
    n_high = sum(1 for x in pv["flags"] if x["severity"] == "Alta")
    cards = [
        (pv["non_official"], "Da fonti non ufficiali", S.AMBER if pv["non_official"] else S.TEAL,
         "git, URL, percorsi locali"),
        (len(pv["flags"]), "Segnali di opacità", S.RED if n_high else (S.AMBER if pv["flags"] else S.TEAL),
         f"{n_high} gravi"),
        (f"{pv['attested']}/{pv['evaluated']}" if pv["evaluated"] else "n/d", "Con attestazione SLSA", S.INK,
         "dipendenze verificate su deps.dev" if pv["evaluated"] else "serve la rete (deps.dev)"),
        (len(pv["unsupported"]), "File non letti", S.AMBER if pv["unsupported"] else S.TEAL,
         "dipendenze fuori dalla SBOM"),
    ]
    w, gap = Inches(2.85), Inches(0.23)
    for i, (val, label, col, sub) in enumerate(cards):
        stat_card(s, Inches(0.6) + i * (w + gap), Inches(1.6), w, Inches(1.6), val, label, col, sub)
    rows, colors = [], {}
    for i, x in enumerate(pv["flags"][:6]):
        rows.append([x["severity"], short(x["title"], 58), short(", ".join(x["items"][:3]), 60)
                     + (f" (+{len(x['items']) - 3})" if len(x["items"]) > 3 else "")])
        colors[(i, 0)] = S.RED if x["severity"] == "Alta" else S.AMBER
    if rows:
        table(s, Inches(0.6), Inches(3.45), Inches(12.15), rows, [0.8, 5.0, 6.3], ["Gravità", "Segnale", "Esempi"],
              row_h=Inches(0.42), font=Pt(11), cell_colors=colors)
    else:
        empty_state(s, Inches(0.6), Inches(3.45), Inches(12.15), Inches(1.6), "Nessuna provenienza opaca",
                    "Le dipendenze arrivano dai registri ufficiali, con hash di integrità e senza codice copiato a mano."
                    if pv["source_declared"] else
                    "La SBOM non dichiara la provenienza dei componenti (Syft e Trivy non la riportano): "
                    "restano validi solo i controlli sul repository e su deps.dev.")
    if pv["unsupported"]:
        text_box(s, Inches(0.6), Inches(6.3), Inches(12.1), Inches(0.6),
                 "Non letti: " + short("; ".join(pv["unsupported"]), 230), size=Pt(11), color=S.MUTED)
    notes(s, "Una dipendenza presa da git, da un URL o da un indice alternativo non passa dai controlli del registro "
             "pubblico: chi controlla quella fonte può cambiarne il contenuto. Un nome che non esiste nel registro "
             "pubblico può essere registrato da un attaccante (dependency confusion). L'hash di integrità nel "
             "lockfile garantisce che venga installato sempre lo stesso contenuto.")


def slide_code(deck, f, meta):
    cd = f["code"]
    s = deck.new_slide("Cosa dice il codice che la SBOM non vede", "Segnali nel codice")
    n_high = sum(1 for x in cd["findings"] if x["severity"] == "Alta")
    cards = [
        (len(cd["weak_crypto"]), "Crittografia debole", S.RED if cd["weak_crypto"] else S.TEAL,
         ", ".join(cd["weak_crypto"])[:30] or "nessun algoritmo superato"),
        (len(cd["secrets"]), "Segreti in configurazione", S.AMBER if cd["secrets"] else S.TEAL,
         "passphrase, password, chiavi"),
        (cd["credentials_in_url"], "Credenziali nell'indirizzo", S.RED if cd["credentials_in_url"] else S.TEAL,
         "in cronologia e nei log"),
        (len(cd["endpoints"]), "Indirizzi esterni", S.AMBER if cd["endpoints"] else S.TEAL,
         "servizi da cui dipende"),
    ]
    w, gap = Inches(2.85), Inches(0.23)
    for i, (val, label, col, sub) in enumerate(cards):
        stat_card(s, Inches(0.6) + i * (w + gap), Inches(1.6), w, Inches(1.6), val, label, col, sub)
    rows, colors = [], {}
    for i, x in enumerate(cd["findings"][:5]):
        rows.append([x["severity"], short(x["title"], 62), short(", ".join(x["where"][:2]), 52)
                     + (f" (+{x['count'] - 2})" if x["count"] > 2 else "")])
        colors[(i, 0)] = S.RED if x["severity"] == "Alta" else S.AMBER
    if rows:
        table(s, Inches(0.6), Inches(3.45), Inches(12.15), rows, [0.8, 5.6, 5.7], ["Gravità", "Segnale", "Dove"],
              row_h=Inches(0.42), font=Pt(11), cell_colors=colors)
    else:
        empty_state(s, Inches(0.6), Inches(3.45), Inches(12.15), Inches(1.6), "Nessun segnale nel codice",
                    "Nessuna crittografia debole, credenziale nell'indirizzo o segreto nei file di configurazione.")
    if cd["endpoints"]:
        text_box(s, Inches(0.6), Inches(6.3), Inches(12.1), Inches(0.6),
                 "Indirizzi esterni: " + short(", ".join(e["host"] for e in cd["endpoints"]), 220),
                 size=Pt(11), color=S.MUTED)
    notes(s, f"{n_high} segnali gravi. Controlli testuali sul codice del progetto, escluso il codice di terzi copiato: "
             "indicano dove guardare, non sostituiscono una revisione del codice. Un segreto letto dal JavaScript "
             "arriva nel browser di ogni utente; 3DES ed ECB sono superati; una password nell'indirizzo resta "
             "nella cronologia e nei log dei server.")


def slide_maintenance(deck, f, meta):
    mt = f["maintenance"]
    s = deck.new_slide("Le dipendenze sono ancora mantenute?", "Manutenibilità")
    cards = [
        (mt["evaluated"], "Dipendenze valutate", S.INK, "dati del registro (deps.dev)"),
        (len(mt["stale"]), "Senza release da 18 mesi", S.AMBER if mt["stale"] else S.TEAL, "probabilmente abbandonate"),
        (len(mt["deprecated"]), "Deprecate", S.RED if mt["deprecated"] else S.TEAL, "dichiarate dall'autore"),
        (len(mt["behind"]), "Indietro di 2+ versioni", S.AMBER if mt["behind"] else S.TEAL, "aggiornamento difficile"),
    ]
    w, gap = Inches(2.85), Inches(0.23)
    for i, (val, label, col, sub) in enumerate(cards):
        stat_card(s, Inches(0.6) + i * (w + gap), Inches(1.6), w, Inches(1.6), val, label, col, sub)
    seen, rows = set(), []
    for kind, items in (("Deprecata", mt["deprecated"]), ("Ferma", mt["stale"]), ("Indietro", mt["behind"])):
        for x in items:
            if x["name"] in seen or len(rows) >= 7:
                continue
            seen.add(x["name"])
            rows.append([short(x["name"], 40), short(x["version"] or "?", 18), short(x["latest"] or "?", 18),
                         f"{x['days']} giorni" if x.get("days") is not None else "n/d", kind])
    if rows:
        table(s, Inches(0.6), Inches(3.45), Inches(12.15), rows, [4, 1.8, 1.8, 2, 1.3],
              ["Componente", "Versione usata", "Ultima versione", "Ultima release", "Stato"], row_h=Inches(0.38),
              font=Pt(11))
    else:
        empty_state(s, Inches(0.6), Inches(3.45), Inches(12.15), Inches(1.6), "Dipendenze in salute",
                    "Nessuna dipendenza valutata risulta abbandonata, deprecata o molto indietro.")
    notes(s, f"Punteggio di manutenibilità: {mt['score']}/100 = 100 − 200 × quota non mantenute o deprecate − "
             "50 × quota indietro di 2 o più versioni maggiori. Le dipendenze dirette vengono interrogate per prime.")


def slide_recommendations(deck, f, meta):
    recs = f["recommendations"][:6]
    s = deck.new_slide("Cosa dovrebbe fare la PA", "Raccomandazioni")
    col_w, gap = Inches(5.95), Inches(0.25)
    card_h = Inches(1.5)
    for i, (prio, title, body) in enumerate(recs):
        col, row = i % 2, i // 2
        x = Inches(0.6) + col * (col_w + gap)
        y = Inches(1.65) + row * (card_h + Inches(0.18))
        rect(s, x, y, col_w, card_h, S.TINT)
        circle_label(s, x + Inches(0.22), y + Inches(0.22), Inches(0.5), i + 1, S.PRIORITY_COLORS[prio])
        text_box(s, x + Inches(0.9), y + Inches(0.16), col_w - Inches(1.1), Inches(0.35), title,
                 size=Pt(15), color=S.INK, bold=True)
        text_box(s, x + Inches(0.9), y + Inches(0.5), col_w - Inches(1.1), Inches(0.3),
                 S.PRIORITY_LABEL[prio].upper(), size=Pt(10), color=S.PRIORITY_COLORS[prio], bold=True)
        text_box(s, x + Inches(0.9), y + Inches(0.76), col_w - Inches(1.1), Inches(0.7), short(body, 200),
                 size=Pt(12), color=S.TEXT)
    notes(s, "\n".join(f"{i + 1}. [{S.PRIORITY_LABEL[p]}] {t}: {b}" for i, (p, t, b) in enumerate(f["recommendations"])))


def _yes_no(v, yes="Sì", no="No", none="n/d"):
    return none if v is None else (yes if v else no)


def slide_release(deck, f, meta):
    rel = meta["release"]
    pkg = meta.get("packages") or {}
    s = deck.new_slide("Il pacchetto che si installa", "Provenienza del rilascio")
    chk, dig, sigs = rel["checksum_match"], rel["github_digest_match"], rel["signatures"]
    col = lambda v: S.MUTED if v is None else (S.TEAL if v else S.RED)
    cards = [
        (_yes_no(chk), "Impronta dell'autore", col(chk),
         f"file {short(rel['declared_from'], 26)}" if rel["declared_from"] else "nessun file di checksum"),
        (_yes_no(dig), "Digest di GitHub", col(dig), "calcolato da GitHub"),
        (len(sigs), "Firme o attestazioni", S.TEAL if sigs else S.RED,
         short(", ".join(sigs), 30) if sigs else "nessuna nel rilascio"),
        (pkg.get("debs", 0), "Pacchetti .deb", S.INK if pkg.get("expanded") == pkg.get("debs") else S.AMBER,
         f"{pkg.get('expanded', 0)} aperti e analizzati" if not pkg.get("skipped") or pkg.get("expanded")
         else "non aperti: vedi note"),
    ]
    w, gap = Inches(2.85), Inches(0.23)
    for i, (val, label, c, sub) in enumerate(cards):
        stat_card(s, Inches(0.6) + i * (w + gap), Inches(1.6), w, Inches(1.6), val, label, c, sub)
    rows = [
        ["File analizzato", short(rel["asset"], 70)],
        ["Dimensione", f"{rel['size'] / 1024 ** 2:.1f} MB" if rel["size"] >= 1024 ** 2 else f"{rel['size'] / 1024:.1f} KB"],
        ["SHA-256 calcolato", rel["sha256"]],
        ["Pubblicato il / da", f"{rel['published_at'][:10]} · {rel['author'] or 'n/d'}"],
        ["Rilascio", short(rel["url"], 70)],
    ]
    table(s, Inches(0.6), Inches(3.45), Inches(7.6), rows, [1.8, 5.8], ["Evidenza", "Valore"],
          row_h=Inches(0.36), font=Pt(10.5))
    rect(s, Inches(8.5), Inches(3.45), Inches(4.25), Inches(2.25), S.TINT)
    text_box(s, Inches(8.75), Inches(3.6), Inches(3.8), Inches(0.3), "Che cosa garantisce",
             size=Pt(14), color=S.INK, bold=True)
    text_box(s, Inches(8.75), Inches(3.98), Inches(3.8), Inches(1.65),
             "Un'impronta pubblicata accanto al file prova che il download non si è corrotto, non "
             "che il file sia autentico: chi controlla l'account può cambiare file e impronta insieme. "
             "Solo una firma del produttore o un'attestazione di provenienza (SLSA) lega il binario "
             "al codice e alla pipeline che lo hanno prodotto.", size=Pt(11.5), color=S.TEXT)
    notes(s, f"SHA-256 completo: {rel['sha256']}. Dichiarato dall'autore: {rel['declared_sha256'] or 'n/d'}. "
             f"Digest GitHub: {rel['github_digest'] or 'n/d'}. Altri file nel rilascio: "
             + ", ".join(rel["other_assets"]) + "."
             + (" Pacchetti non aperti: " + "; ".join(pkg.get("skipped", [])[:5]) if pkg.get("skipped") else ""))


def slide_governance(deck, f, meta):
    gv = f["governance"]
    s = deck.new_slide("Chi mantiene il software, e con quali regole", "Governance")
    sig = gv["signals"]
    # i segnali sono una ventina: l'altezza delle righe si adatta allo spazio disponibile (1.6"-6.9")
    row_h = min(Inches(0.4), Emu(int(Inches(5.3) / max(1, len(sig)))))
    dot = Emu(int(row_h * 0.72))
    label_size = Pt(13) if len(sig) <= 13 else Pt(10.5)
    for i, x in enumerate(sig):
        y = Inches(1.6) + i * row_h
        sym, col = ("✓", S.TEAL) if x["ok"] else (("✗", S.RED) if x["ok"] is False else ("–", S.MUTED))
        circle_label(s, Inches(0.6), y, dot, sym, col, size=Pt(8) if len(sig) > 13 else Pt(10))
        text_box(s, Inches(0.98), y + Emu(int(row_h * 0.08)), Inches(4.45), row_h, short(x["label"], 58),
                 size=label_size, color=S.TEXT if x["ok"] is not None else S.MUTED)
        if x["note"]:
            text_box(s, Inches(5.45), y + Emu(int(row_h * 0.12)), Inches(2.65), row_h, short(x["note"], 44),
                     size=Pt(9), color=S.MUTED)
    hist, sc = gv.get("history"), gv.get("scorecard")
    # al massimo tre schede, in ordine di importanza: concentrazione, Scorecard, ultimo commit, punteggio
    cards = []
    if hist:
        share = round(100 * hist["top_author_share"])
        cards.append((f"{share}%", "dei commit a un solo autore", S.RED if share > 60 else S.TEAL,
                      f"{hist['authors']} autori negli ultimi {hist['commits']} commit"))
    if sc:
        bp = gv.get("best_practices")
        cards.append((f"{sc['score']}/10", "OpenSSF Scorecard", S.level_color(round(sc["score"] * 10)),
                      f"branch principale, {sc['date']}" + (f" · badge {bp}" if bp else "")))
    if gv.get("last_commit_days") is not None:
        d = gv["last_commit_days"]
        cards.append((f"{d} gg", "dall'ultimo commit", S.TEAL if d <= 183 else S.AMBER,
                      f"ultimo tag: {short(gv['last_tag'], 36)}" if gv.get("last_tag") else "nessun tag"))
    cards = cards[:2] + [(f"{gv['score']}/100", "Segnali soddisfatti", S.level_color(gv["score"]),
                          "tra quelli applicabili")]
    for i, (val, label, col, sub) in enumerate(cards):
        stat_card(s, Inches(8.25), Inches(1.65) + i * Inches(1.72), Inches(4.5), Inches(1.58), val, label, col, sub)
    ex = gv.get("exit") or {}
    notes(s, "Segnali letti dai file del repository, dallo storico git e (in rete) dalle API di GitHub, "
             "Scorecard e OpenSSF Best Practices: nessun dato è dichiarato a mano. "
             "Il simbolo – indica un segnale non applicabile o non verificabile (non entra nel punteggio). "
             "Strategia di uscita: licenza OSI e documentazione su esportazione, migrazione e backup dei dati "
             f"(trovati: {', '.join(ex.get('signals') or []) or 'nessun riferimento'}). "
             "La concentrazione dei commit misura il 'bus factor': se una persona fa quasi tutto, "
             "la continuità del software dipende da lei."
             + (" OpenSSF Scorecard si riferisce al branch principale, non necessariamente alla versione analizzata."
                if sc else ""))


def slide_pipeline(deck, f, meta):
    pl = f["pipeline"]
    s = deck.new_slide("Come nasce il binario", "Filiera di build")
    unp = pl["actions_total"] - pl["actions_pinned"]
    cards = [
        (pl["pipe_to_shell"], "Esecuzioni 'curl | bash'", S.RED if pl["pipe_to_shell"] else S.TEAL,
         "eseguito senza verifica"),
        (f"{unp}/{pl['actions_total']}", "Azioni non bloccate", S.AMBER if unp else S.TEAL,
         f"{pl['actions_third_party_unpinned']} di autori esterni"),
        (f"{pl['images_unpinned']}/{pl['images_total']}", "Immagini senza digest",
         S.AMBER if pl["images_unpinned"] else S.TEAL,
         f"{pl.get('images_eol', 0)} su sistemi fuori supporto"),
        ("Sì" if pl["has_signing"] else "No", "Rilasci firmati", S.TEAL if pl["has_signing"] else S.RED,
         f"checksum: {'sì' if pl['has_checksum'] else 'no'} · SBOM: {'sì' if pl['has_sbom_step'] else 'no'}"),
    ]
    w, gap = Inches(2.85), Inches(0.23)
    for i, (val, label, col, sub) in enumerate(cards):
        stat_card(s, Inches(0.6) + i * (w + gap), Inches(1.6), w, Inches(1.6), val, label, col, sub)
    groups = {}
    for x in pl["findings"]:
        key = (x["severity"], x["kind"], x["text"])
        g = groups.setdefault(key, {"n": 0, "where": x["where"], "detail": x["detail"]})
        g["n"] += 1
    rows, colors = [], {}
    for i, ((sev, kind, text), g) in enumerate(list(groups.items())[:6]):
        rows.append([sev, kind, short(text, 44) + (f"  (×{g['n']})" if g["n"] > 1 else ""), short(g["where"], 54)])
        colors[(i, 0)] = S.RED if sev == "Alta" else S.AMBER
    if rows:
        table(s, Inches(0.6), Inches(3.45), Inches(12.15), rows, [0.75, 1.75, 4.3, 4.4],
              ["Gravità", "Problema", "Elemento", "Primo punto trovato"], row_h=Inches(0.4), font=Pt(11),
              cell_colors=colors)
    else:
        empty_state(s, Inches(0.6), Inches(3.45), Inches(12.15), Inches(1.6), "Nessun problema nella pipeline",
                    "Azioni e immagini sono bloccate per impronta e non ci sono download eseguiti senza verifica.")
    more = len(pl["findings"]) - sum(g["n"] for g in list(groups.values())[:6])
    text_box(s, Inches(0.6), Inches(6.5), Inches(12.1), Inches(0.4),
             (f"Altri {more} punti nel file risultati.json. " if more > 0 else "") +
             "Un tag (@v4) è un'etichetta che il proprietario può spostare; un hash o un digest sha256 no.",
             size=S.SMALL_SIZE, color=S.MUTED)
    notes(s, "La SBOM dice di cosa è fatto il software; questa slide dice se il processo che lo produce è "
             "verificabile. Esempio reale: nel 2025 l'azione tj-actions/changed-files è stata compromessa e i "
             "suoi tag spostati su codice malevolo; chi la usava con un hash di commit non è stato colpito. "
             f"Punteggio della filiera: {pl['score']}/100.")


def slide_method(deck, f, meta):
    s = deck.new_slide("Metodo ed evidenze verificabili", "Metodologia")
    n_dims = sum(1 for v in f["sovereignty"]["dimensions"].values() if v >= 0)
    steps = [("Raccolta", "SBOM generata (Syft, Trivy o motore integrato) o fornita"),
             ("Analisi", "Vulnerabilità + KEV + EPSS, licenze, provenienza, manutenzione, governance, build"),
             ("Indice", f"{n_dims} dimensioni con formule pubbliche"),
             ("Evidenze", "Input, risultati e report con impronta SHA-256")]
    for i, (t, d) in enumerate(steps):
        y = Inches(1.7) + i * Inches(1.05)
        circle_label(s, Inches(0.6), y, Inches(0.55), i + 1, S.INK)
        text_box(s, Inches(1.35), y, Inches(4.3), Inches(0.35), t, size=Pt(16), color=S.INK, bold=True)
        text_box(s, Inches(1.35), y + Inches(0.38), Inches(4.3), Inches(0.55), d, size=S.SMALL_SIZE, color=S.MUTED)
    eng = meta["engines"]
    git = meta.get("git") or {}
    rows = [["Sorgente analizzata", short(git.get("url") or meta["source"], 46)]]
    if git.get("commit"):
        rows.append(["Riferimento / commit", short(f"{git.get('ref', '')} · {git['commit'][:12]}", 46)])
    if meta.get("release"):
        rows.append(["File del rilascio", short(meta["release"]["asset"], 46)])
    if meta.get("archive_sha256"):
        rows.append(["SHA-256 dell'archivio", meta["archive_sha256"][:16] + "…" + meta["archive_sha256"][-8:]])
    rows += [
        ["SBOM", short(f"{meta['sbom_format']} · {meta['sbom_generator']}", 46)],
        ["SHA-256 della SBOM", meta["sbom_sha256"][:16] + "…" + meta["sbom_sha256"][-8:]],
        ["Motore vulnerabilità", short(eng["vulns"], 46)],
        ["Database vulnerabilità", short(eng["vulns_db"], 46)],
        ["KEV / EPSS", short(f"{meta['enrichment'].get('kev', 'non consultato')} / "
                             f"{meta['enrichment'].get('epss', 'non consultato')}", 46)],
    ]
    table(s, Inches(6.0), Inches(1.7), Inches(6.75), rows, [2.2, 4.3], ["Evidenza", "Valore"],
          row_h=Inches(0.34), font=Pt(11))
    y = Inches(1.7) + Inches(0.34) * (len(rows) + 1) + Inches(0.25)
    rect(s, Inches(6.0), y, Inches(6.75), Inches(1.4), S.TINT)
    text_box(s, Inches(6.25), y + Inches(0.12), Inches(6.3), Inches(0.3), "Limiti dichiarati",
             size=Pt(14), color=S.INK, bold=True)
    text_box(s, Inches(6.25), y + Inches(0.45), Inches(6.3), Inches(0.9),
             "La SBOM descrive i componenti, non il comportamento: dipendenze di rete a runtime e "
             "giurisdizione dei fornitori richiedono analisi aggiuntive. I risultati valgono alla data "
             "del database vulnerabilità.", size=Pt(12), color=S.TEXT)
    online = meta.get("online_sources") or []
    notes(s, "Tutti i file sono nella cartella del report con SHA256SUMS: chiunque può verificare con "
             "'sha256sum -c SHA256SUMS' e rifare l'analisi con gli stessi strumenti.\n"
             f"SHA-256 completo della SBOM: {meta['sbom_sha256']}\n"
             + (f"Commit completo: {git['commit']}\n" if git.get("commit") else "")
             + ("Servizi online consultati: " + ", ".join(online) + "\n" if online else
                online_note(meta) + "\n")
             + ("Comandi eseguiti:\n" + "\n".join(meta["commands"]) if meta["commands"] else ""))


# ================================================================== entry
def build_presentation(findings: dict, meta: dict, out_path: str) -> str:
    deck = Deck(meta)
    has_comps = findings["composition"]["total"] > 0
    slide_title(deck, findings, meta)
    slide_summary(deck, findings, meta)
    slide_sovereignty(deck, findings, meta)
    if meta.get("release"):
        slide_release(deck, findings, meta)
    if has_comps:
        slide_composition(deck, findings, meta)
        slide_quality(deck, findings, meta)
        if findings["vulns"] is not None:
            slide_vulns(deck, findings, meta)
            slide_priorities(deck, findings, meta)
            slide_exposed(deck, findings, meta)
        slide_licenses(deck, findings, meta)
        if findings.get("provenance", {}).get("applicable"):
            slide_provenance(deck, findings, meta)
        if findings.get("maintenance", {}).get("applicable"):
            slide_maintenance(deck, findings, meta)
        slide_duplicates(deck, findings, meta)
    if findings["vendor_diff"] is not None:
        slide_vendor(deck, findings, meta)
    if not has_comps and findings.get("provenance", {}).get("applicable"):
        slide_provenance(deck, findings, meta)
    if findings.get("code", {}).get("applicable"):
        slide_code(deck, findings, meta)
    if findings.get("governance", {}).get("applicable"):
        slide_governance(deck, findings, meta)
    if findings.get("pipeline", {}).get("applicable"):
        slide_pipeline(deck, findings, meta)
    slide_recommendations(deck, findings, meta)
    slide_method(deck, findings, meta)
    deck.prs.save(out_path)
    return out_path
